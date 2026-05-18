"""Parse uploaded business documents into plain text for AI knowledge."""

from __future__ import annotations

import csv
import io
from pathlib import Path


class DocumentParseError(ValueError):
    pass


def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return "pdf"
    if ext in (".xlsx", ".xls"):
        return "xlsx"
    if ext == ".csv":
        return "csv"
    raise DocumentParseError(f"Unsupported file type: {ext or '(no extension)'}")


def parse_document(file_path: Path, file_type: str) -> str:
    if file_type == "pdf":
        return _parse_pdf(file_path)
    if file_type == "xlsx":
        return _parse_xlsx(file_path)
    if file_type == "csv":
        return _parse_csv(file_path)
    raise DocumentParseError(f"Unknown file type: {file_type}")


def _parse_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise DocumentParseError("PDF parser not installed") from e

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text.strip())
    if not parts:
        raise DocumentParseError("No text could be extracted from this PDF")
    return "\n\n".join(parts)


def _parse_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise DocumentParseError("Excel parser not installed") from e

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"--- Sheet: {sheet.title} ---\n" + "\n".join(rows))
    wb.close()
    if not parts:
        raise DocumentParseError("No data found in this spreadsheet")
    return "\n\n".join(parts)


def _parse_csv(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="replace")
    reader = csv.reader(io.StringIO(raw))
    rows: list[str] = []
    for row in reader:
        cells = [c.strip() for c in row if c and c.strip()]
        if cells:
            rows.append(" | ".join(cells))
    if not rows:
        raise DocumentParseError("No data found in this CSV file")
    return "\n".join(rows)
